from __future__ import annotations

import asyncio
import datetime as dt
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Dict, List, Optional, Sequence, Tuple
from urllib.parse import urlparse

import httpx
from docx import Document
from pypdf import PdfReader

from .config import Settings
from .utils import html_to_text, iso, utc_now


class CanvasError(RuntimeError):
    pass


@dataclass
class CanvasSource:
    source_key: str
    course_id: str
    kind: str
    title: str
    url: str
    text: str
    updated_at: Optional[str]


class CanvasClient:
    def __init__(self, token: str, settings: Settings, transport: Optional[httpx.AsyncBaseTransport] = None):
        self.settings = settings
        self.base_url = settings.canvas_base_url.rstrip("/")
        self._host = urlparse(self.base_url).netloc
        self.client = httpx.AsyncClient(
            base_url=self.base_url,
            headers={
                "Authorization": "Bearer " + token,
                "Accept": "application/json",
                "User-Agent": "EmberHacks-SchoolSync/1.0",
            },
            timeout=httpx.Timeout(45.0, connect=15.0),
            follow_redirects=False,
            transport=transport,
        )

    async def __aenter__(self) -> "CanvasClient":
        return self

    async def __aexit__(self, *_args: Any) -> None:
        await self.client.aclose()

    async def _request(self, path_or_url: str, params: Optional[Sequence[Tuple[str, str]]] = None) -> httpx.Response:
        if path_or_url.startswith("http"):
            parsed = urlparse(path_or_url)
            if parsed.scheme != "https" or parsed.netloc != self._host or not parsed.path.startswith("/api/v1/"):
                raise CanvasError("Canvas returned an unsafe pagination URL")
            url = path_or_url
        else:
            if not path_or_url.startswith("/api/v1/"):
                raise CanvasError("Canvas API path must start with /api/v1/")
            url = path_or_url
        for attempt in range(4):
            try:
                response = await self.client.get(url, params=params if attempt == 0 else None)
            except (httpx.TimeoutException, httpx.NetworkError) as exc:
                if attempt == 3:
                    raise CanvasError("Quercus is temporarily unreachable") from exc
                await asyncio.sleep(2**attempt)
                continue
            if response.status_code in {429, 500, 502, 503, 504} and attempt < 3:
                retry_after = response.headers.get("Retry-After", "")
                await asyncio.sleep(min(int(retry_after), 60) if retry_after.isdigit() else 2**attempt)
                continue
            if response.status_code == 401:
                raise CanvasError("The Quercus API token is invalid or expired")
            if response.status_code == 403:
                raise CanvasError("Quercus denied access to a requested course resource")
            if response.status_code == 404:
                raise CanvasError("Quercus resource was not found")
            try:
                response.raise_for_status()
            except httpx.HTTPStatusError as exc:
                raise CanvasError("Quercus API returned HTTP %s" % response.status_code) from exc
            return response
        raise CanvasError("Quercus request failed")

    async def get_object(self, path: str, params: Optional[Sequence[Tuple[str, str]]] = None) -> Dict[str, Any]:
        response = await self._request(path, params)
        value = response.json()
        if not isinstance(value, dict):
            raise CanvasError("Quercus returned an unexpected object response")
        return value

    async def paginate(self, path: str, params: Optional[Sequence[Tuple[str, str]]] = None) -> List[Dict[str, Any]]:
        query = list(params or []) + [("per_page", "100")]
        url: Optional[str] = path
        rows: List[Dict[str, Any]] = []
        seen = set()
        while url:
            if url in seen or len(seen) >= 500:
                raise CanvasError("Quercus returned a pagination loop")
            seen.add(url)
            response = await self._request(url, query if len(seen) == 1 else None)
            value = response.json()
            if not isinstance(value, list):
                raise CanvasError("Quercus returned an unexpected list response")
            rows.extend(row for row in value if isinstance(row, dict))
            url = self._next_url(response.headers.get("Link", ""))
        return rows

    @staticmethod
    def _next_url(header: str) -> Optional[str]:
        for part in header.split(","):
            match = re.search(r'<([^>]+)>;\s*rel="next"', part)
            if match:
                return match.group(1)
        return None

    async def current_user(self) -> Dict[str, Any]:
        return await self.get_object("/api/v1/users/self")

    async def courses(self) -> List[Dict[str, Any]]:
        return await self.paginate("/api/v1/courses", [
            ("enrollment_type", "student"),
            ("include[]", "term"),
            ("include[]", "enrollments"),
            ("include[]", "sections"),
            ("include[]", "syllabus_body"),
            ("state[]", "available"),
            ("state[]", "completed"),
        ])

    async def assignments(self, course_id: str) -> List[Dict[str, Any]]:
        return await self.paginate("/api/v1/courses/%s/assignments" % course_id, [
            ("include[]", "submission"), ("include[]", "rubric"),
            ("include[]", "all_dates"), ("order_by", "due_at"),
        ])

    async def announcements(self, course_id: str) -> List[Dict[str, Any]]:
        now = utc_now()
        return await self.paginate("/api/v1/announcements", [
            ("context_codes[]", "course_" + course_id),
            ("start_date", iso(now - dt.timedelta(days=180)) or ""),
            ("end_date", iso(now + dt.timedelta(days=30)) or ""),
            ("active_only", "true"),
        ])

    async def calendar_events(self, course_id: str) -> List[Dict[str, Any]]:
        year = self.settings.academic_term_year
        if self.settings.academic_term_season == "fall":
            start = dt.datetime(year, 8, 1, tzinfo=dt.timezone.utc)
            end = dt.datetime(year + 1, 1, 15, 23, 59, tzinfo=dt.timezone.utc)
        else:
            start = dt.datetime(year, 1, 1, tzinfo=dt.timezone.utc)
            end = dt.datetime(year + 1, 1, 1, tzinfo=dt.timezone.utc)
        rows = await self.paginate("/api/v1/calendar_events", [
            ("type", "event"), ("context_codes[]", "course_" + course_id),
            ("start_date", iso(start) or ""),
            ("end_date", iso(end) or ""),
            ("includes[]", "series_natural_language"),
        ])
        child_ids = sorted({
            str(child)
            for row in rows if row.get("hidden")
            for child in row.get("child_events") or []
            if isinstance(child, (int, str))
        })
        if child_ids:
            values = await asyncio.gather(
                *(self.get_object("/api/v1/calendar_events/" + child_id) for child_id in child_ids),
                return_exceptions=True,
            )
            details = {
                child_id: value for child_id, value in zip(child_ids, values)
                if isinstance(value, dict)
            }
            for row in rows:
                row["child_events"] = [
                    details.get(str(child), child) if isinstance(child, (int, str)) else child
                    for child in row.get("child_events") or []
                ]
        return rows

    async def planner_items(self, course_id: str) -> List[Dict[str, Any]]:
        now = utc_now()
        return await self.paginate("/api/v1/planner/items", [
            ("context_codes[]", "course_" + course_id),
            ("start_date", iso(now - dt.timedelta(days=30)) or ""),
            ("end_date", iso(now + dt.timedelta(days=365)) or ""),
        ])

    async def course_sources(self, course: Dict[str, Any]) -> List[CanvasSource]:
        course_id = str(course["id"])
        base = self.base_url + "/courses/" + course_id
        sources: List[CanvasSource] = []
        syllabus = html_to_text(course.get("syllabus_body"))
        if syllabus:
            sources.append(CanvasSource("syllabus:" + course_id, course_id, "syllabus",
                                        "Course syllabus", base + "/assignments/syllabus", syllabus,
                                        course.get("updated_at")))

        try:
            front = await self.get_object("/api/v1/courses/%s/front_page" % course_id)
            text = html_to_text(front.get("body"))
            if text:
                sources.append(CanvasSource("front_page:" + course_id, course_id, "front_page",
                                            front.get("title") or "Course home page", base,
                                            text, front.get("updated_at")))
        except CanvasError:
            pass

        try:
            pages = await self.paginate("/api/v1/courses/%s/pages" % course_id,
                                        [("published", "true"), ("include[]", "body")])
            for page in pages[:500]:
                text = html_to_text(page.get("body"))
                if text:
                    sources.append(CanvasSource(
                        "page:%s:%s" % (course_id, page.get("page_id") or page.get("url")),
                        course_id, "page", page.get("title") or "Course page",
                        base + "/pages/" + str(page.get("url", "")), text, page.get("updated_at")))
        except CanvasError:
            pass

        try:
            files = await self.paginate("/api/v1/courses/%s/files" % course_id,
                                        [("sort", "updated_at"), ("order", "desc")])
            candidates = [row for row in files if self._likely_schedule_file(row)][:30]
            extracted = await asyncio.gather(*(self._extract_file(course_id, row) for row in candidates))
            sources.extend(row for row in extracted if row is not None)
        except CanvasError:
            pass

        try:
            modules = await self.paginate("/api/v1/courses/%s/modules" % course_id)
            module_text = []
            for module in modules[:100]:
                items = await self.paginate("/api/v1/courses/%s/modules/%s/items" % (course_id, module["id"]))
                module_text.append(module.get("name", "Module"))
                module_text.extend(str(item.get("title", "")) for item in items[:200])
            if module_text:
                sources.append(CanvasSource("modules:" + course_id, course_id, "modules",
                                            "Course modules", base + "/modules", "\n".join(module_text), None))
        except CanvasError:
            pass
        return sources

    @staticmethod
    def _likely_schedule_file(row: Dict[str, Any]) -> bool:
        name = str(row.get("display_name") or row.get("filename") or "").casefold()
        content_type = str(row.get("content-type") or row.get("content_type") or "")
        relevant = re.search(r"syllabus|course.?outline|schedule|calendar|assessment|assignment", name)
        supported = content_type in {
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "text/plain", "text/markdown", "text/csv",
        } or name.endswith((".pdf", ".docx", ".txt", ".md", ".csv"))
        return bool(relevant and supported and int(row.get("size") or 0) <= 20 * 1024 * 1024)

    async def _extract_file(self, course_id: str, row: Dict[str, Any]) -> Optional[CanvasSource]:
        url = row.get("url")
        if not url:
            return None
        parsed = urlparse(url)
        if parsed.scheme != "https":
            return None
        async with httpx.AsyncClient(timeout=60, follow_redirects=True) as client:
            async with client.stream("GET", url) as response:
                response.raise_for_status()
                if response.url.scheme != "https":
                    raise CanvasError("Course file redirected to an unsafe URL")
                data = bytearray()
                async for chunk in response.aiter_bytes():
                    data.extend(chunk)
                    if len(data) > self.settings.max_download_bytes:
                        raise CanvasError("Course file exceeded the download limit")
        name = str(row.get("display_name") or row.get("filename") or "Course file")
        try:
            text = extract_document(bytes(data), name)
        except Exception:
            return None
        if not text:
            return None
        return CanvasSource("file:%s:%s" % (course_id, row["id"]), course_id, "file", name,
                            self.base_url + "/courses/%s/files/%s" % (course_id, row["id"]),
                            text, row.get("updated_at"))


def extract_document(data: bytes, name: str) -> str:
    lower = name.casefold()
    if lower.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)[:100_000]
    if lower.endswith(".docx"):
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)[:100_000]
    return data.decode("utf-8", errors="replace")[:100_000]
